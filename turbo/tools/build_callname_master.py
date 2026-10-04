#!/usr/bin/env python3
"""Build an FC 27 callname master (workbook + Turbo masters JSON) for one commentary language.

Port of turbo_dev\\masters\\build_master_v0.py (2026-10-04, ita_it), including its later column H "nameid" and the
Play macro patch; the JSON also carries the audio folder and the segments per id (Turbo's play buttons).

Inputs:
  --bank        <lang>_bank.json, extracted from the FC 27 commentary files on disk by turbo/tools/fc27_commentary
                (keys: lang, game_build, real [{segment, variation, playerid}], real_link [...], generic
                [{segment, variation, commentaryid}])
  --players     turbo_table_players.csv (FC 27 career database export: Turbo Tools > Exports > Export tables)
  --edited      turbo_table_editedplayernames.csv (same export)
  --names       bridge_names.txt (name id<TAB>text; FC 27's playernames table is empty in the career database)
  --commentary  bridge_commentary.txt (commentary id<TAB>text)
  --template    the user's FC 26 workbook of that language (e.g. C:\\FC_Tools\\My Mods\\ita\\italy_master.xlsm). It is
                only read and copied, never written: its styles, column widths, other sheets and VBA project are kept.
                Its 'callnames' rows also give names and categories (column G) for ids FC 27 has no name for, and its
                'real' rows of players in the FC 27 database teach the nationality -> category map.
  --lang        language code for the JSON (default: the bank's "lang")
  --wav-dir     folder holding real\\ and generic\\ (the segment wavs: real\\pPLAYER_NAMES_SIMPLE_<seg>_<seg>.wav,
                generic\\pSIMPLE_SURNAME_<seg>_<seg>.wav), written to the JSON for Turbo's play buttons; every row's wav
                is checked there (default: the bank's "wav_dir", where fc27_commentary wrote them). Italian:
                C:\\FC_Tools\\My Mods\\i27
  --vba-from / --vba-to
                same-length text patch of the copied workbook's VBA project (vba_tool.patch_text: module source,
                p-code and __SRP caches), e.g. the Play macro's audio base C:\\FC_Tools\\My Mods\\ita\\ ->
                C:\\FC_Tools\\My Mods\\i27\\ (the FC 27 audio). Without them the VBA project stays byte-identical.
  --copy-to     also copy the finished workbook into this folder (build_master_v0 delivers it to the FC 27 audio folder)

Outputs:
  --out-xlsm    <name>_master_fc27.xlsm: sheet 'callnames' (SegmentID, VariationId, playerid, name, type, Play, category,
                nameid) one row per recording, sorted accent-insensitively by name (then real before generic, id,
                segment, variation; rows without a name last); H 'nameid' (generic rows) = the playernames id(s) whose
                text equals the commentary name (exact text, else case-insensitive; several joined with ','): the id a
                player's last / common name must be set to for that recording. Sheet 'names' (id, name). The VBA
                project byte-identical unless --vba-from/--vba-to. When the file is open in Excel, <name>_new.xlsm is
                written instead.
  --out-json    <lang>.json for Turbo (turbo\\callnames\\masters):
                {"turbo_masters": 1, "language", "game": "fc27", "game_build", "source", "built", "counts": {...},
                 "real_players" (PLAYER_NAMES_SIMPLE | PLAYER_NAMES_LINK ids = players with their own recording),
                 "real_simple_players", "real_link_players", "generic_ids", "names" {"<playerid>": name, SIMPLE only},
                 "generic_names" {"<commentaryid>": text}, "wav_dir" (--wav-dir),
                 "segments" {"generic": {"<commentaryid>": [seg, ...]}, "real": {"<playerid>": [seg, ...]} (SIMPLE),
                             "real_link": {...} only when wav_dir holds pPLAYER_NAMES_LINK wavs (only those segments)}}

Audio: the user's Play macro hard-codes C:\\FC_Tools\\My Mods\\<lang folder>\\ as the audio base (the FC 26 folders).
The FC 27 workbook is shipped with its own real\\ and generic\\ folders of FC 27 recordings next to it (--vba-from /
--vba-to point its macro there); this tool does not write audio.

usage: python turbo/tools/build_callname_master.py --bank B --players P --edited E --names N --commentary C
                                                    --template T --out-xlsm X --out-json J [--lang ita_it]
                                                    [--wav-dir W] [--vba-from OLD --vba-to NEW] [--copy-to D]
       python turbo/tools/build_callname_master.py --self-test
ita_it as build_master_v0: --wav-dir "C:\\FC_Tools\\My Mods\\i27" --vba-from "C:\\FC_Tools\\My Mods\\ita\\"
                           --vba-to "C:\\FC_Tools\\My Mods\\i27\\" --copy-to "C:\\FC_Tools\\My Mods\\i27"
"""
import argparse
import collections
import csv
import datetime
import json
import os
import shutil
import sys
import tempfile
import unicodedata
import zipfile
from copy import copy

PLAY = '\u25b6 Play'
WAV_PREFIX = {'real': 'pPLAYER_NAMES_SIMPLE', 'real_link': 'pPLAYER_NAMES_LINK', 'generic': 'pSIMPLE_SURNAME'}
WAV_SUB = {'real': 'real', 'real_link': 'real', 'generic': 'generic'}


def wav_path(wav_dir, kind, seg):
    """<wav_dir>\\real|generic\\<prefix>_<seg>_<seg>.wav (the names fc27_commentary and the Play macro use)"""
    return os.path.join(wav_dir, WAV_SUB[kind], f'{WAV_PREFIX[kind]}_{seg}_{seg}.wav')


def read_tsv_map(path):
    m = {}
    with open(path, encoding='utf-8', errors='replace') as f:  # some names are cut inside a UTF-8 sequence
        for line in f:
            if line.startswith('#'):
                continue
            a = line.rstrip('\n').split('\t')
            if len(a) >= 2 and a[0].strip().isdigit():
                m[int(a[0])] = a[1]
    return m


def sort_key(s):
    s = unicodedata.normalize('NFKD', s or '')
    return ''.join(ch for ch in s if not unicodedata.combining(ch)).casefold()


def read_csv_by_playerid(path):
    out = {}
    with open(path, encoding='utf-8', newline='') as f:
        for r in csv.DictReader(f):
            out[int(r['playerid'])] = r
    return out


def vba(path):
    with zipfile.ZipFile(path) as z:
        return z.read('xl/vbaProject.bin')


def nameid_index(names):
    """name text -> name ids, exact and case-insensitive (column H)"""
    exact, folded = collections.defaultdict(list), collections.defaultdict(list)
    for nid, txt in names.items():
        exact[txt.strip()].append(nid)
        folded[txt.strip().casefold()].append(nid)
    return exact, folded


def nameids_for(text, exact, folded, stats):
    """Column H: the name id (an int), several joined with ',' (a str), None when no name has this text"""
    t = (text or '').strip()
    ids = exact.get(t)
    if ids:
        stats['exact'] += 1
    else:
        ids = folded.get(t.casefold())
        stats['casefold' if ids else 'none'] += 1
    if not ids:
        return None
    ids = sorted(ids)
    return ids[0] if len(ids) == 1 else ','.join(map(str, ids))


def segments_by_id(rows, key):
    out = collections.defaultdict(set)
    for b in rows:
        out[int(b[key])].add(int(b['segment']))
    return {str(i): sorted(v) for i, v in sorted(out.items())}


def build(bank_path, players_path, edited_path, names_path, comm_path, template, out_xlsm, out_json, lang=None,
          wav_dir=None, vba_from=None, vba_to=None, copy_to=None, log=print):
    import openpyxl

    with open(bank_path, encoding='utf-8') as f:
        bank = json.load(f)
    lang = lang or bank.get('lang')
    names = read_tsv_map(names_path)
    comm = read_tsv_map(comm_path)
    players = read_csv_by_playerid(players_path)
    edited = read_csv_by_playerid(edited_path)

    def player_name(pid):
        e = edited.get(pid)
        if e:
            common = (e.get('commonname') or '').strip()
            if common:
                return common
            return ' '.join(x for x in ((e.get('firstname') or '').strip(), (e.get('surname') or '').strip()) if x)
        p = players.get(pid)
        if not p:
            return None

        def nm(field):
            v = int(p.get(field) or 0)
            return names.get(v, '').strip() if v > 0 else ''

        common = nm('commonnameid')
        if common:
            return common
        return ' '.join(x for x in (nm('firstnameid'), nm('lastnameid')) if x)

    # FC 26 template rows: names and categories by (type, id)
    wb26 = openpyxl.load_workbook(template, read_only=True, data_only=True)
    cat26, name26 = {}, {}
    for row in wb26['callnames'].iter_rows(min_row=2, values_only=True):
        if len(row) < 7 or row[4] not in ('real', 'generic') or row[2] is None:
            continue
        key = (row[4], int(row[2]))
        if row[6]:
            cat26.setdefault(key, row[6])
        if row[3] and not str(row[3]).startswith('='):
            name26.setdefault(key, str(row[3]))
    wb26.close()

    # nationality -> category, learned from the FC 26 'real' rows of players that are in the FC 27 database
    votes = collections.defaultdict(collections.Counter)
    for (typ, pid), cat in cat26.items():
        if typ == 'real' and pid in players:
            votes[players[pid].get('nationality')][cat] += 1
    nat_cat = {n: c.most_common(1)[0][0] for n, c in votes.items()}

    rows = []
    missing_names = collections.Counter()
    for typ, key, src in (('real', 'playerid', bank['real']), ('generic', 'commentaryid', bank['generic'])):
        for b in src:
            i = int(b[key])
            if typ == 'real':
                n = player_name(i) or name26.get(('real', i))
                cat = cat26.get(('real', i)) or (nat_cat.get(players[i].get('nationality')) if i in players else None)
            else:
                n = comm.get(i) or name26.get(('generic', i))
                cat = cat26.get(('generic', i))
            if not n:
                missing_names[typ] += 1
                n = ''
            rows.append((int(b['segment']), int(b['variation']), i, n, typ, cat))
    rows.sort(key=lambda r: (sort_key(r[3]) or '~', r[4] != 'real', r[2], r[0], r[1]))

    # column H 'nameid' (generic rows): the playernames id(s) whose text equals the commentary name
    exact, folded = nameid_index(names)
    nameid_stats = collections.Counter()

    # workbook: copy the template (to a temporary name), replace the data rows, keep styles, widths and the VBA project
    out_dir = os.path.dirname(os.path.abspath(out_xlsm))
    os.makedirs(out_dir, exist_ok=True)
    building = out_xlsm + '.building.xlsm'
    shutil.copyfile(template, building)
    wb = openpyxl.load_workbook(building, keep_vba=True)
    ws = wb['callnames']
    style_row = [copy(c._style) for c in ws[3]]   # a plain 'real' row of the template
    head_style = copy(ws.cell(row=1, column=7)._style)
    ws.delete_rows(2, ws.max_row)
    for r_i, (seg, var, i, n, typ, cat) in enumerate(rows, start=2):
        nid = nameids_for(n, exact, folded, nameid_stats) if typ == 'generic' else None
        for c_i, v in enumerate((seg, var, i, n, typ, PLAY, cat, nid), start=1):
            c = ws.cell(row=r_i, column=c_i, value=v)
            c._style = copy(style_row[c_i - 1] if c_i != 8 else style_row[2])  # H styled like C (an id)
    h = ws.cell(row=1, column=8, value='nameid')
    h._style = head_style
    ws.column_dimensions['H'].width = max(ws.column_dimensions['C'].width or 12, 14)
    ws.auto_filter.ref = f'A1:H{len(rows) + 1}'
    ws.freeze_panes = 'A2'

    wn = wb['names']
    wn.delete_rows(2, wn.max_row)
    idnames = {}
    for seg, var, i, n, typ, cat in rows:
        if n:
            idnames.setdefault(i, n)
    for r_i, (i, n) in enumerate(sorted(idnames.items()), start=2):
        wn.cell(row=r_i, column=1, value=i)
        wn.cell(row=r_i, column=2, value=n)
    wb.save(building)

    if vba(building) != vba(template):
        os.remove(building)
        raise RuntimeError('VBA project changed by openpyxl')
    if vba_from or vba_to:
        sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
        import vba_tool  # same folder; stdlib only (OLE + MS-OVBA)
        vba_new, report = vba_tool.patch_text(vba(template), vba_from or '', vba_to or '')
        vba_tool.replace_zip_member(building, 'xl/vbaProject.bin', vba_new)
        if vba(building) != vba_new:
            raise RuntimeError('VBA project not replaced')
        log('VBA patched: ' + '; '.join(report))
    try:
        os.replace(building, out_xlsm)
    except PermissionError:   # the previous build is open in Excel
        out_xlsm = out_xlsm[:-5] + '_new.xlsm'
        os.replace(building, out_xlsm)
        log('NOTE: the previous workbook is open in Excel; wrote ' + out_xlsm)
    if copy_to:
        shutil.copyfile(out_xlsm, os.path.join(copy_to, os.path.basename(out_xlsm)))
        log('copied to ' + copy_to)
    log('nameid (generic rows): ' + str(dict(nameid_stats)))

    # audio: the folder Turbo's play buttons read (every row's wav is checked there) and the segments of every id
    wav_dir = wav_dir if wav_dir is not None else (bank.get('wav_dir') or '')
    if wav_dir and os.path.isdir(wav_dir):
        miss = [(typ, seg, i) for seg, var, i, n, typ, cat in rows if not os.path.isfile(wav_path(wav_dir, typ, seg))]
        log(f'wav_dir {wav_dir}: rows without a wav: {len(miss)} {miss[:10]}')
    elif wav_dir:
        log(f'wav_dir {wav_dir} is not a folder here: the wavs were not checked')
    segments = {'generic': segments_by_id(bank['generic'], 'commentaryid'), 'real': segments_by_id(bank['real'], 'playerid')}
    link = [b for b in bank.get('real_link', [])
            if wav_dir and os.path.isfile(wav_path(wav_dir, 'real_link', int(b['segment'])))]
    if link:
        segments['real_link'] = segments_by_id(link, 'playerid')
    real_ids = sorted({int(b['playerid']) for b in bank['real']})
    link_ids = sorted({int(b['playerid']) for b in bank.get('real_link', [])})
    gen_ids = sorted({int(b['commentaryid']) for b in bank['generic']})
    real_set = set(real_ids)
    counts = {'rows': len(rows), 'real_rows': sum(r[4] == 'real' for r in rows),
              'generic_rows': sum(r[4] == 'generic' for r in rows), 'real_players': len(real_ids),
              'real_link_players': len(link_ids), 'own_recording_players': len(real_set | set(link_ids)),
              'generic_ids': len(gen_ids), 'real_in_fc27_db': sum(1 for i in real_ids if i in players),
              'missing_names': dict(missing_names), 'category_empty': sum(1 for r in rows if not r[5])}
    out = {'turbo_masters': 1, 'language': lang, 'game': 'fc27', 'game_build': bank.get('game_build'),
           'source': out_xlsm, 'built': datetime.datetime.now().isoformat(timespec='seconds'), 'counts': counts,
           'real_players': sorted(real_set | set(link_ids)), 'real_simple_players': real_ids,
           'real_link_players': link_ids, 'generic_ids': gen_ids,
           'names': {str(i): n for i, n in sorted(idnames.items()) if i in real_set},
           'generic_names': {str(i): comm.get(i, '') for i in gen_ids},
           'wav_dir': wav_dir, 'segments': segments}
    out_json_dir = os.path.dirname(os.path.abspath(out_json))
    os.makedirs(out_json_dir, exist_ok=True)
    with open(out_json, 'w', encoding='utf-8') as f:
        json.dump(out, f, ensure_ascii=False, indent=1)
    return out


# ---------------------------------------------------------------------------------------------------------- self-test

HEADER = ('SegmentID', 'VariationId', 'playerid', 'name', 'type', 'Play Audio', 'category')
DUMMY_VBA = b'\xd0\xcf\x11\xe0dummy vba project \x00\x01\x02' * 7


def _make_template(path):
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'callnames'
    ws.append(HEADER)
    ws.append((1, 10, 900001, 'Rossi', 'generic', PLAY, 'Generic'))
    ws.append((2, 20, 100, 'Old Name', 'real', PLAY, 'Serie A'))   # row 3: the style row
    ws.append((3, 30, 200, 'Zeta', 'real', PLAY, 'Liga'))
    for c in ws[3]:
        c.font = Font(bold=True, color='FF123456')
        c.fill = PatternFill('solid', fgColor='FFFFEE00')
    ws.column_dimensions['D'].width = 33
    wn = wb.create_sheet('names')
    wn.append(('id', 'name'))
    wn.append((1, 'stale'))
    wb.create_sheet('fc26 heads').append(('kept',))
    plain = path + '.tmp.xlsx'
    wb.save(plain)
    with zipfile.ZipFile(plain) as zin, zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == '[Content_Types].xml':
                data = data.replace(b'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml',
                                    b'application/vnd.ms-excel.sheet.macroEnabled.main+xml')
                data = data.replace(b'</Types>', b'<Default Extension="bin" '
                                    b'ContentType="application/vnd.ms-office.vbaProject"/></Types>')
            if item.filename == 'xl/_rels/workbook.xml.rels':
                data = data.replace(b'</Relationships>', b'<Relationship Id="rIdVBA" Type="http://schemas.microsoft.'
                                    b'com/office/2006/relationships/vbaProject" Target="vbaProject.bin"/>'
                                    b'</Relationships>')
            zout.writestr(item, data)
        zout.writestr('xl/vbaProject.bin', DUMMY_VBA)
    os.remove(plain)


def self_test():
    import openpyxl

    fails = []

    def check(cond, what):
        print(('  ok    ' if cond else '  FAIL  ') + what)
        if not cond:
            fails.append(what)

    with tempfile.TemporaryDirectory() as d:
        p = lambda n: os.path.join(d, n)  # noqa: E731
        template = p('test_master.xlsm')
        _make_template(template)
        bank = {'turbo_bank': 1, 'lang': 'ita_it', 'game_build': '1.0.test',
                'real': [{'segment': 5, 'variation': 2, 'playerid': 100}, {'segment': 5, 'variation': 1, 'playerid': 100},
                         {'segment': 6, 'variation': 1, 'playerid': 200}, {'segment': 7, 'variation': 1, 'playerid': 300},
                         {'segment': 8, 'variation': 1, 'playerid': 400}],
                'real_link': [{'segment': 9, 'variation': 1, 'playerid': 100},
                              {'segment': 10, 'variation': 2, 'playerid': 500}],
                'generic': [{'segment': 1, 'variation': 10, 'commentaryid': 900001},
                            {'segment': 2, 'variation': 11, 'commentaryid': 900002}]}
        with open(p('ita_it_bank.json'), 'w', encoding='utf-8') as f:
            json.dump(bank, f)
        with open(p('players.csv'), 'w', encoding='utf-8', newline='') as f:
            f.write('playerid,firstnameid,lastnameid,commonnameid,nationality\n'
                    '100,1,2,0,27\n200,3,4,0,45\n300,0,0,5,27\n')
        with open(p('edited.csv'), 'w', encoding='utf-8', newline='') as f:
            f.write('playerid,firstname,surname,commonname\n200,Ángel,Álvarez,\n')
        with open(p('names.txt'), 'w', encoding='utf-8') as f:
            # 6 / 9 'Rossi': two exact matches for column H; 7 'rossi' loses to them; 8 'ÄBEL' matches 'Äbel' by case
            f.write('# nameid\tname\n1\tÉmile\n2\tZola\n3\tX\n4\tY\n5\tBeta\n6\tRossi\n7\trossi\n8\tÄBEL\n9\tRossi\n')
        with open(p('comm.txt'), 'w', encoding='utf-8') as f:
            f.write('900001\tRossi\n900002\tÄbel\n')
        wav_dir = p('wavs')
        for sub, name in (('generic', 'pSIMPLE_SURNAME_1_1.wav'), ('real', 'pPLAYER_NAMES_SIMPLE_5_5.wav'),
                          ('real', 'pPLAYER_NAMES_LINK_9_9.wav')):   # the LINK wav of segment 10 is missing
            os.makedirs(os.path.join(wav_dir, sub), exist_ok=True)
            open(os.path.join(wav_dir, sub, name), 'wb').close()
        out_xlsm, out_json = p(os.path.join('out', 'test_master_fc27.xlsm')), p(os.path.join('out', 'ita_it.json'))
        logged = []
        out = build(p('ita_it_bank.json'), p('players.csv'), p('edited.csv'), p('names.txt'), p('comm.txt'),
                    template, out_xlsm, out_json, wav_dir=wav_dir, log=logged.append)

        wb = openpyxl.load_workbook(out_xlsm, keep_vba=True)
        ws = wb['callnames']
        got = [tuple(c.value for c in r) for r in ws.iter_rows(min_row=1, max_row=ws.max_row)]
        want = [HEADER + ('nameid',),
                (2, 11, 900002, 'Äbel', 'generic', PLAY, None, 8),
                (6, 1, 200, 'Ángel Álvarez', 'real', PLAY, 'Liga', None),
                (7, 1, 300, 'Beta', 'real', PLAY, 'Serie A', None),
                (5, 1, 100, 'Émile Zola', 'real', PLAY, 'Serie A', None),
                (5, 2, 100, 'Émile Zola', 'real', PLAY, 'Serie A', None),
                (1, 10, 900001, 'Rossi', 'generic', PLAY, 'Generic', '6,9'),
                (8, 1, 400, None, 'real', PLAY, None, None)]   # '' is written as an empty cell
        check(got[0] == want[0], 'header row kept, H = nameid')
        check(len(got) == len(want), f'rows: {len(got) - 1} data rows (want {len(want) - 1})')
        check(got == want, 'sort order (accent-insensitive by name, nameless last), names, categories, nameid (generic)')
        if got != want:
            for g in got:
                print('        ', g)
        check(all(c.font.bold and c.font.color.rgb == 'FF123456' and c.fill.fgColor.rgb == 'FFFFEE00'
                  for r in ws.iter_rows(min_row=2) for c in r), 'styles copied from the template style row')
        check(ws.column_dimensions['D'].width == 33, 'column width kept')
        check(ws.auto_filter.ref == f'A1:H{len(want)}' and ws.freeze_panes == 'A2', 'filter and frozen header')
        check(ws.cell(row=2, column=8).font.bold and ws.column_dimensions['H'].width >= 14, 'column H styled and wide')
        check(not os.path.exists(out_xlsm + '.building.xlsm'), 'no temporary workbook left')
        wn = wb['names']
        check([tuple(c.value for c in r) for r in wn.iter_rows(min_row=2)] ==
              [(100, 'Émile Zola'), (200, 'Ángel Álvarez'), (300, 'Beta'), (900001, 'Rossi'), (900002, 'Äbel')],
              'names sheet')
        check('fc26 heads' in wb.sheetnames, 'other sheets kept')
        check(vba(out_xlsm) == DUMMY_VBA == vba(template), 'VBA bytes identical')

        with open(out_json, encoding='utf-8') as f:
            j = json.load(f)
        check(list(j) == ['turbo_masters', 'language', 'game', 'game_build', 'source', 'built', 'counts',
                          'real_players', 'real_simple_players', 'real_link_players', 'generic_ids', 'names',
                          'generic_names', 'wav_dir', 'segments'], 'json keys')
        check(j['wav_dir'] == wav_dir, 'json wav_dir')
        check(j['segments'] == {'generic': {'900001': [1], '900002': [2]},
                                'real': {'100': [5], '200': [6], '300': [7], '400': [8]},
                                'real_link': {'100': [9]}}, 'json segments (LINK only where its wav exists)')
        check(any('rows without a wav: 4' in s for s in logged), 'missing wavs counted: ' + '; '.join(logged))
        check(j['turbo_masters'] == 1 and j['language'] == 'ita_it' and j['game'] == 'fc27'
              and j['game_build'] == '1.0.test' and j['source'] == out_xlsm, 'json header fields')
        check(j['real_players'] == [100, 200, 300, 400, 500] and j['real_simple_players'] == [100, 200, 300, 400]
              and j['real_link_players'] == [100, 500] and j['generic_ids'] == [900001, 900002], 'json id lists')
        check(j['names'] == {'100': 'Émile Zola', '200': 'Ángel Álvarez', '300': 'Beta'}, 'json names (SIMPLE only)')
        check(j['generic_names'] == {'900001': 'Rossi', '900002': 'Äbel'}, 'json generic names')
        check(j['counts'] == {'rows': 7, 'real_rows': 5, 'generic_rows': 2, 'real_players': 4, 'real_link_players': 2,
                              'own_recording_players': 5, 'generic_ids': 2, 'real_in_fc27_db': 3,
                              'missing_names': {'real': 1}, 'category_empty': 2}, 'json counts')
        check(out == j, 'returned dict equals the written JSON')
        # no --wav-dir and none in the bank: an empty wav_dir, the SIMPLE / generic segments, no real_link
        out2 = build(p('ita_it_bank.json'), p('players.csv'), p('edited.csv'), p('names.txt'), p('comm.txt'),
                     template, p(os.path.join('out2', 'm.xlsm')), p(os.path.join('out2', 'ita_it.json')), log=logged.append)
        check(out2['wav_dir'] == '' and sorted(out2['segments']) == ['generic', 'real'], 'no wav folder: segments only')
    print('SELF-TEST ' + ('PASS' if not fails else f'FAIL ({len(fails)})'))
    return 0 if not fails else 1


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split('\n')[0])
    ap.add_argument('--bank')
    ap.add_argument('--players')
    ap.add_argument('--edited')
    ap.add_argument('--names')
    ap.add_argument('--commentary')
    ap.add_argument('--template')
    ap.add_argument('--out-xlsm')
    ap.add_argument('--out-json')
    ap.add_argument('--lang')
    ap.add_argument('--wav-dir')
    ap.add_argument('--vba-from')
    ap.add_argument('--vba-to')
    ap.add_argument('--copy-to')
    ap.add_argument('--self-test', action='store_true')
    a = ap.parse_args(argv)
    if a.self_test:
        return self_test()
    need = ('bank', 'players', 'edited', 'names', 'commentary', 'template', 'out_xlsm', 'out_json')
    missing = ['--' + k.replace('_', '-') for k in need if not getattr(a, k)]
    if missing:
        ap.error('missing ' + ', '.join(missing))
    if os.path.abspath(a.out_xlsm) == os.path.abspath(a.template):
        ap.error('--out-xlsm must not be the template (the template is never written)')
    if bool(a.vba_from) != bool(a.vba_to):
        ap.error('--vba-from and --vba-to go together')
    out = build(a.bank, a.players, a.edited, a.names, a.commentary, a.template, a.out_xlsm, a.out_json, a.lang,
                wav_dir=a.wav_dir, vba_from=a.vba_from, vba_to=a.vba_to, copy_to=a.copy_to)
    print(json.dumps(out['counts'], indent=1))
    print(f"wrote {out['source']}\nwrote {a.out_json}")
    return 0


if __name__ == '__main__':
    sys.exit(main())
