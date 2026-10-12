#!/usr/bin/env python3
"""Import the user's callname master workbooks into Turbo's per-language lists (turbo\\callnames\\masters).

The user mapped, for FC 26, every recording of each commentary language in one Excel workbook per language
(<root>\\<folder>\\<name>_master.xlsm, sheet "callnames", header row SegmentID, VariationId, playerid, name, type, ...).
An FC 27 master built from the game's own data uses the same workbook format and is named <name>_master_fc27.xlsm
(or .xlsx); when a language has both, the FC 27 master is used:

  type 'real'     the recording belongs to ONE player: the playerid column is that player's id. The game says these
                  players' names from their own recordings (PLAYER_LOW_SIMPLE / PLAYER_LOW_LINK with player_db_pID),
                  whatever playernamemap or the name ids say (docs/callnames.md sections 1 and 5.3).
  type 'generic'  a generic surname recording: the playerid column holds its COMMENTARY id (9xxxxx).
An FC 27 master written by build_callname_master.py also holds the bank's other families (by player id unless said):
  type 'real link'     pPLAYER_NAMES_LINK: also an own recording; some players have LINK rows only
  type 'real high'     pPLAYER_NAMES_HIGH by player id (not an own recording for the game's kick-off check)
  type 'generic high'  pPLAYER_NAMES_HIGH by surname id (commentary id; none in FC 27 1.0.140 ita_it)
Any other type is counted (other_type_rows) and skipped.

FC 27 mostly reuses the FC 26 recordings, so without an FC 27 master the FC 26 lists are the best available answer to
"has this player his own recording?" - the game's audio service only answers that during a match, and its set is
incomplete. The Turbo GUI (Players > Callname) reads <Live Editor>\\turbo\\callnames\\masters\\<lang>.json and shows
the source as "FC 27 master" or "your FC 26 list" (the "game" field below).

usage: python turbo/tools/import_callname_masters.py [--root DIR ...] [--le "C:\\FC 27 Live Editor"]
                                                       [--out DIR] [--lang ita_it ...] [--dry-run] [--self-test]
                                                       [--wav-dir LANG=DIR ...] [--game-build LANG=BUILD ...]
  --root (repeatable) defaults to "C:\\FC_Tools\\My Mods" and "<le>\\turbo_dev\\masters" (each when present).
  --out defaults to <le>\\turbo\\callnames\\masters. The workbooks are only read (openpyxl, read-only mode).
  --wav-dir dut_nl=C:\\FC_Tools\\My Mods\\ned\\audio   one language's audio folder (the folder holding generic\\ and
         real\\); "dut_nl=" writes none. Without it (resolve_wav_dir): an FC 27 Italian master -> <root>\\i27, else the
         workbook's folder, else <workbook folder>\\audio (ned), else <root>\\<first 3 letters of the code>[\\audio].
  --game-build ita_it=1.0.140.64835   the game build an FC 27 master was extracted from. Without it
         (find_game_build): the JSON build_callname_master.py wrote for the same workbook (<lang>.json in a root, beside
         or above the workbook), else the bank <root>\\raw\\<lang>_bank.json; none found = no "game_build" key.

Output, one file per language (<out>\\<lang>.json, UTF-8):
  {"turbo_masters": 1, "language": "ita_it", "game": "fc27" | "fc26", ["game_build": "<FC 27 build>",]
   "source": "<workbook path>", "built": "<ISO time>", "note": "...", "counts": {...},
   "real_players": [sorted unique player ids of 'real' and 'real link' rows = the players with their own recording],
   ["real_simple_players": ['real' ids], "real_link_players": ['real link' ids],]
   "generic_ids": [sorted unique commentary ids of 'generic' rows],
   ["real_high_players": ['real high' ids], "generic_high_ids": ['generic high' ids],]
   "names": {"<playerid>": "<name>", ...} (real_players: the name of the 'real' rows, else of the 'real link' rows),
   "generic_names": {"<commentaryid>": "<name>"}, ["wav_dir": "<folder holding generic\\ and real\\>",]
   "segments": {"generic": {"<id>": [SegmentID, ...]}, "real": {...}, ["real_link", "real_high", "generic_high"]}}
  The bracketed FC 27 keys are written only when the workbook has rows of those types, so an FC 26 list keeps its
  shape; they are named as in build_callname_master.py's JSON (the GUI's loader, callnames.cpp, reads both).
"""
import argparse
import collections
import datetime
import json
import os
import sys
import tempfile

# Folder under the root -> FC 27 commentary language code (the pack names under <game>\commentary and Data\Win32:
# commentaryfull_<code>). The workbook's own name is the fallback for a file found elsewhere.
FOLDER_LANG = {"br": "por_br", "eng": "eng_us", "fra": "fre_fr", "ger": "ger_de", "ita": "ita_it", "ned": "dut_nl",
               "spa": "spa_es"}
STEM_LANG = {"br": "por_br", "uk": "eng_us", "eng": "eng_us", "france": "fre_fr", "fra": "fre_fr", "ger": "ger_de",
             "italy": "ita_it", "ita": "ita_it", "dutch": "dut_nl", "ned": "dut_nl", "spain": "spa_es", "spa": "spa_es"}
EXT_RANK = {".xlsm": 0, ".xlsx": 1}  # several workbooks for one language: .xlsm first (ned has an older .xlsx copy)
# Which game a workbook describes, by its name: <stem>_master_fc27.xls[mx] = FC 27 (built from the game), <stem>_master
# = the user's FC 26 list. An FC 27 master wins over an FC 26 list of the same language.
SUFFIXES = (("_master_fc27", "fc27"), ("_master", "fc26"))
GAME_RANK = {"fc27": 0, "fc26": 1}
GAME_LABEL = {"fc27": "FC 27 master", "fc26": "FC 26 list"}
SHEET = "callnames"
COMMENTARY_MIN, COMMENTARY_MAX = 900001, 965000  # the player-name range every commentary bank uses
PLAYER_ID_MAX = 999999                            # FC player ids are below a million; bigger ones are bank oddities
# Workbook type (column E, any case, '_' = ' ') -> family key in the JSON. 'real' / 'generic' are the FC 26 types; the
# others are the FC 27 bank families build_callname_master.py adds (its KINDS). Commentary-id families: generic*.
FAMILIES = {"real": "real", "generic": "generic", "real link": "real_link", "real high": "real_high",
            "generic high": "generic_high"}
BY_COMMENTARY_ID = ("generic", "generic_high")
FC27_FAMILIES = ("real_link", "real_high", "generic_high")
# the FC 27 Italian audio fc27_commentary extracted (2026-10-04): <root>\i27\{generic,real,real_link,real_high}
FC27_AUDIO_FOLDER = {"ita_it": "i27"}


def default_le_root():
    return os.environ.get("TURBO_LE_ROOT", r"C:\FC 27 Live Editor")


def default_roots(le):
    """The user's FC 26 workbooks, then the folder the FC 27 masters are built into."""
    return [r"C:\FC_Tools\My Mods", os.path.join(le, "turbo_dev", "masters")]


def workbook_kind(file_name):
    """(name without the suffix, "fc27" | "fc26") for a master workbook's file name, else None."""
    stem, ext = os.path.splitext(file_name)
    if file_name.startswith("~$") or ext.lower() not in EXT_RANK:
        return None
    low = stem.lower()
    for suffix, game in SUFFIXES:
        if low.endswith(suffix):
            return low[: -len(suffix)], game
    return None


def find_workbooks(roots):
    """{lang: [(path, game), best first]} for every *_master.xls[mx] / *_master_fc27.xls[mx] under each root (one
    folder level deep, plus the root itself). Best first: an FC 27 master, then .xlsm, then the newest."""
    if isinstance(roots, str):
        roots = [roots]
    found = collections.defaultdict(list)
    for root in roots:
        if not os.path.isdir(root):
            continue
        dirs = [root] + [os.path.join(root, d) for d in sorted(os.listdir(root)) if os.path.isdir(os.path.join(root, d))]
        for d in dirs:
            folder = os.path.basename(d).lower() if d != root else ""
            for f in sorted(os.listdir(d)):
                kind = workbook_kind(f)
                if not kind:
                    continue
                base, game = kind
                lang = FOLDER_LANG.get(folder) or STEM_LANG.get(base)
                if lang:
                    found[lang].append((os.path.join(d, f), game))
    for books in found.values():
        books.sort(key=lambda b: (GAME_RANK[b[1]], EXT_RANK[os.path.splitext(b[0])[1].lower()], -os.path.getmtime(b[0])))
    return dict(found)


def as_id(v, lowest=1):
    """An id cell as an int >= lowest, else None (text, blanks, fractions, smaller numbers). Ids are positive; a
    SegmentID may be 0 (lowest=0: the bank's first segment, pSIMPLE_SURNAME_0_0.wav exists in i27\\generic)."""
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, int):
        return v if v >= lowest else None
    if isinstance(v, float):
        return int(v) if v.is_integer() and v >= lowest else None
    s = str(v).strip()
    if s.isdigit():
        n = int(s)
        return n if n >= lowest else None
    return None


def family_of(type_cell):
    """The JSON family key of a workbook type cell ('Real ', 'real link', 'Real_High' ...), else None."""
    return FAMILIES.get(" ".join(str(type_cell or "").replace("_", " ").split()).lower())


def read_workbook(path):
    """Rows of the "callnames" sheet -> (families, counts). families[key] = {"ids": {id: Counter of names} (player
    families) or {id: first name} (commentary families), "segments": {id: [SegmentID, ...]}} for every key of
    FAMILIES. Raises ValueError when the workbook has no such sheet or header."""
    import openpyxl  # imported here so --help works without it

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = next((n for n in wb.sheetnames if n.strip().lower() == SHEET), None)
        if sheet is None:
            raise ValueError(f"no sheet '{SHEET}' (sheets: {', '.join(wb.sheetnames)})")
        rows = wb[sheet].iter_rows(values_only=True)
        header = next(rows, None)
        cols = {str(h).strip().lower(): i for i, h in enumerate(header or ()) if h is not None}
        missing = [c for c in ("playerid", "name", "type") if c not in cols]
        if missing:
            raise ValueError(f"sheet '{sheet}' has no column {', '.join(missing)} in its header row")
        cs = cols.get("segmentid")
        ci, cn, ct = cols["playerid"], cols["name"], cols["type"]
        fams = {key: {"ids": {} if key in BY_COMMENTARY_ID else collections.defaultdict(collections.Counter),
                      "segments": collections.defaultdict(list)} for key in FAMILIES.values()}
        first_seen = {}
        c = collections.Counter()
        for row in rows:
            if row is None or all(v is None for v in row):
                c["blank_rows"] += 1
                continue
            c["rows"] += 1
            get = lambda i: row[i] if i < len(row) else None  # noqa: E731 - short rows end early in openpyxl
            key = family_of(get(ct))
            if key is None:
                c["other_type_rows"] += 1
                continue
            pid = as_id(get(ci))
            if pid is None:
                c["non_numeric_id_rows"] += 1
                continue
            seg = as_id(get(cs), lowest=0) if cs is not None else None
            name = str(get(cn) or "").strip()
            fam = fams[key]
            c[key + "_rows"] += 1
            if key in BY_COMMENTARY_ID:
                if name and not fam["ids"].get(pid):
                    fam["ids"][pid] = name
                else:
                    fam["ids"].setdefault(pid, "")
            elif name:
                fam["ids"][pid][name] += 1
                first_seen.setdefault((key, pid, name), c["rows"])
            else:
                fam["ids"][pid]  # noqa: B018 - the id counts even without a name
            if seg is not None:
                fam["segments"][pid].append(seg)
        for key, fam in fams.items():
            if key not in BY_COMMENTARY_ID:
                # the most frequent spelling; ties go to the first one in the sheet
                fam["names"] = {pid: min(counter.items(), key=lambda kv: (-kv[1], first_seen[(key, pid, kv[0])]))[0]
                                for pid, counter in fam["ids"].items() if counter}
        return fams, c
    finally:
        wb.close()


def parse_lang_map(items, what):
    """["ita_it=VALUE", ...] -> {"ita_it": "VALUE"}; an empty VALUE is kept (it means 'none'). ValueError otherwise."""
    out = {}
    for it in items or ():
        lang, sep, value = str(it).partition("=")
        lang = lang.strip().lower()
        if not sep or len(lang) != 6 or lang[3] != "_":
            raise ValueError(f"{what} '{it}': expected <lang>=<value>, e.g. dut_nl=C:\\FC_Tools\\My Mods\\ned\\audio")
        out[lang] = value.strip().strip('"')
    return out


def has_audio(folder):
    """True when the folder holds a generic\\ or real\\ subfolder (any case: br has Generic\\ and Real\\)."""
    try:
        with os.scandir(folder) as it:
            return any(e.is_dir() and e.name.lower() in ("generic", "real") for e in it)
    except OSError:
        return False


def resolve_wav_dir(lang, path, game, roots=(), overrides=None):
    """(audio folder or None, where it came from) for the play buttons (<wav_dir>\\generic, real, real_link ...).
    1. overrides[lang] (--wav-dir; '' = none); 2. an FC 27 master of a language with an extracted FC 27 audio folder:
    <root>\\i27 (ita_it); 3. the workbook's folder; 4. <workbook folder>\\audio (ned keeps its wavs in
    ned\\audio\\{generic,real}); 5. FC 26 lists only: <root>\\<first 3 letters of the code>, then its audio\\ - each
    when it holds a generic\\ or real\\ folder. An FC 27 master never gets an FC 26 audio folder (FC 27 renumbered the
    segments, docs/re/fc27-commentary-bank.md)."""
    if overrides and lang in overrides:
        d = overrides[lang]
        if not d:
            return None, f"none (--wav-dir {lang}=)"
        d = os.path.abspath(d)
        return d, "--wav-dir" + ("" if has_audio(d) else " (WARNING: no generic\\ or real\\ folder there)")
    wb_dir = os.path.dirname(os.path.abspath(path))
    cands = []
    if game == "fc27" and lang in FC27_AUDIO_FOLDER:
        cands += [(os.path.join(r, FC27_AUDIO_FOLDER[lang]), "the FC 27 audio folder") for r in roots]
    cands += [(wb_dir, "the workbook's folder"), (os.path.join(wb_dir, "audio"), "the workbook's audio folder")]
    if game != "fc27":
        for r in roots:
            cands += [(os.path.join(r, lang[:3]), "the language folder"),
                      (os.path.join(r, lang[:3], "audio"), "the language folder's audio")]
    for d, why in cands:
        if has_audio(d):
            return d, why
    return None, "no folder with generic\\ or real\\ found (--wav-dir sets one)"


def _read_json(path):
    try:
        with open(path, encoding="utf-8") as f:
            return json.load(f)
    except (OSError, ValueError):
        return None


def find_game_build(lang, path, roots=(), overrides=None):
    """(FC 27 build or None, where it came from) for an FC 27 master: overrides[lang] (--game-build), else the JSON
    build_callname_master.py wrote for this workbook (<lang>.json beside or above it, or in a root; same "language",
    "game": "fc27" and the same workbook file name in "source"), else the bank fc27_commentary wrote
    (<lang>_bank.json in raw\\ beside or above the workbook, or in <root>\\raw)."""
    if overrides and overrides.get(lang):
        return overrides[lang], "--game-build"
    wb_dir = os.path.dirname(os.path.abspath(path))
    up = os.path.dirname(wb_dir)
    bases = list(dict.fromkeys([wb_dir, up] + [os.path.abspath(r) for r in roots]))
    name = os.path.basename(path).lower()
    for b in bases:
        p = os.path.join(b, f"{lang}.json")
        j = _read_json(p)
        if (isinstance(j, dict) and j.get("game") == "fc27" and j.get("game_build") and str(j.get("language", "")).lower() == lang
                and os.path.basename(str(j.get("source", "")).replace("\\", os.sep)).lower() == name):
            return str(j["game_build"]), p
    for b in bases:
        p = os.path.join(b, "raw", f"{lang}_bank.json")
        j = _read_json(p)
        if isinstance(j, dict) and j.get("game_build") and str(j.get("lang", "")).lower() == lang:
            return str(j["game_build"]), p
    return None, "not found (--game-build sets it)"


def build_list(lang, path, game="fc26", built=None, roots=(), wav_dirs=None, game_builds=None, log=None):
    fams, c = read_workbook(path)
    ids = {k: sorted(f["ids"]) for k, f in fams.items()}
    simple, link = set(ids["real"]), set(ids["real_link"])
    real_ids = sorted(simple | link)   # the players with their own recording: SIMPLE | LINK (build_callname_master)
    generic_ids = ids["generic"]
    fc27_types = any(c[k + "_rows"] for k in FC27_FAMILIES)
    counts = {
        "rows": c["rows"], "blank_rows": c["blank_rows"], "real_rows": c["real_rows"], "generic_rows": c["generic_rows"],
        "other_type_rows": c["other_type_rows"], "non_numeric_id_rows": c["non_numeric_id_rows"],
        "real_players": len(ids["real"]), "generic_ids": len(generic_ids),
        "real_ids_over_999999": sum(1 for p in real_ids if p > PLAYER_ID_MAX),
        "generic_ids_outside_900001_965000": sum(1 for g in generic_ids if not COMMENTARY_MIN <= g <= COMMENTARY_MAX),
    }
    if fc27_types:   # the names build_callname_master.py gives these counts
        counts.update({"real_link_rows": c["real_link_rows"], "real_high_rows": c["real_high_rows"],
                       "generic_high_rows": c["generic_high_rows"], "real_link_players": len(link),
                       "own_recording_players": len(real_ids), "real_high_players": len(ids["real_high"]),
                       "generic_high_ids": len(ids["generic_high"]),
                       "any_recording_players": len(simple | link | set(ids["real_high"]))})
    kinds = ("'real' = the player's own recording (bound to his player id); 'generic' = a generic surname recording "
             "(commentary id).")
    if fc27_types:
        kinds += (" 'real link' = his LINK recordings (also own recordings: real_players = real | real link); "
                  "'real high' / 'generic high' = the HIGH recordings by player id / commentary id.")

    wav_dir, wav_why = resolve_wav_dir(lang, path, game, roots, wav_dirs)
    build, build_why = find_game_build(lang, path, roots, game_builds) if game == "fc27" else (None, "")
    if log:
        log(f"{lang}: wav_dir {wav_dir or 'none'} ({wav_why})" +
            (f"; game_build {build or 'none'} ({build_why})" if game == "fc27" else ""))

    def seg_map(key):
        segs = fams[key]["segments"]
        return {str(i): sorted(dict.fromkeys(segs[i])) for i in ids[key] if i in segs}

    segments = {"generic": seg_map("generic"), "real": seg_map("real")}
    for key in FC27_FAMILIES:
        m = seg_map(key)
        if m:
            segments[key] = m
    simple_names, link_names = fams["real"]["names"], fams["real_link"]["names"]
    names = {}
    for p in real_ids:
        n = simple_names.get(p) or link_names.get(p)
        if n:
            names[str(p)] = n

    res = {"turbo_masters": 1, "language": lang, "game": game}
    if build:
        res["game_build"] = build
    res.update({
        "source": os.path.abspath(path),
        "built": built or datetime.datetime.now().astimezone().isoformat(timespec="seconds"),
        "note": ("FC 27 master: built from FC 27's own data. " if game == "fc27" else
                 "FC 26 data: FC 27 mostly reuses these recordings. ") + kinds,
        "counts": counts,
        "real_players": real_ids,
    })
    if fc27_types:
        res["real_simple_players"] = ids["real"]
        res["real_link_players"] = ids["real_link"]
    res["generic_ids"] = generic_ids
    if fc27_types:
        res["real_high_players"] = ids["real_high"]
        res["generic_high_ids"] = ids["generic_high"]
    res["names"] = names
    gnames = fams["generic"]["ids"]
    res["generic_names"] = {str(g): gnames[g] for g in generic_ids if gnames.get(g)}
    if wav_dir:
        res["wav_dir"] = wav_dir
    if any(segments.values()):
        res["segments"] = {k: v for k, v in segments.items() if v or k in ("generic", "real")}
    return res


def write_json(path, data):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    fd, tmp = tempfile.mkstemp(prefix=".tmp_", suffix=".json", dir=os.path.dirname(path))
    try:
        with os.fdopen(fd, "w", encoding="utf-8", newline="\n") as f:
            json.dump(data, f, ensure_ascii=False, indent=1)
            f.write("\n")
        os.replace(tmp, path)
    except BaseException:
        if os.path.exists(tmp):
            os.remove(tmp)
        raise


def run(roots, out, langs=None, dry_run=False, log=print, wav_dirs=None, game_builds=None):
    if isinstance(roots, str):
        roots = [roots]
    present = [r for r in roots if os.path.isdir(r)]
    for r in roots:
        if r not in present:
            log(f"note: no folder {r} (skipped)")
    if not present:
        log("error: none of the folders exists: " + ", ".join(roots))
        return 2, {}
    books = find_workbooks(present)
    if langs:
        books = {k: v for k, v in books.items() if k in langs}
    if not books:
        log("error: no *_master.xlsm / .xlsx or *_master_fc27.xlsm / .xlsx workbook under " + ", ".join(present))
        return 1, {}
    results, failed = {}, 0
    for lang in sorted(books):
        path, game = books[lang][0]
        for other, other_game in books[lang][1:]:
            log(f"{lang}: also found {other} ({GAME_LABEL[other_game]}): skipped ({os.path.basename(path)} preferred: "
                f"an FC 27 master first, then .xlsm, then the newest)")
        try:
            data = build_list(lang, path, game, roots=present, wav_dirs=wav_dirs, game_builds=game_builds, log=log)
        except Exception as e:  # one unreadable workbook must not stop the others
            failed += 1
            log(f"{lang}: {path}: NOT imported: {e}")
            continue
        results[lang] = data
        k = data["counts"]
        if "own_recording_players" in k:   # an FC 27 master with the bank's other families
            own = (f"{k['own_recording_players']} players with their own recording ({k['real_players']} 'real', "
                   f"{k['real_rows']} rows | {k['real_link_players']} 'real link', {k['real_link_rows']} rows), "
                   f"{k['real_high_players']} players with 'real high' recordings ({k['real_high_rows']} rows), "
                   f"{k['generic_high_ids']} 'generic high' ids ({k['generic_high_rows']} rows)")
        else:
            own = f"{k['real_players']} players with their own recording ('real', {k['real_rows']} rows)"
        log(f"{lang}: {path} ({GAME_LABEL[game]}): {own}, {k['generic_ids']} generic commentary ids ({k['generic_rows']} rows); skipped "
            f"{k['non_numeric_id_rows']} rows with a non-numeric id, {k['other_type_rows']} rows of another type, "
            f"{k['blank_rows']} blank rows; notes: {k['real_ids_over_999999']} real ids over 999999, "
            f"{k['generic_ids_outside_900001_965000']} generic ids outside 900001..965000 (kept)")
        if not dry_run:
            dest = os.path.join(out, f"{lang}.json")
            write_json(dest, data)
            log(f"{lang}: wrote {dest}")
    return (1 if failed and not results else 0), results


def self_test():
    """Synthetic workbooks: language mapping, the .xlsm/.xlsx duplicate, non-numeric ids, names, the JSON written, an
    FC 27 master preferred over the FC 26 list of its language, the FC 27 types ('real link' / 'real high' / 'generic
    high'), game_build, and the audio folder (workbook folder, <folder>\\audio as ned, i27, --wav-dir)."""
    import openpyxl

    with tempfile.TemporaryDirectory() as tmp:
        root, root27, out = os.path.join(tmp, "My Mods"), os.path.join(tmp, "masters"), os.path.join(tmp, "out")

        def book(base, rel, rows, sheet=SHEET, header=("SegmentID", "VariationId", "playerid", "name", "type", "Play Audio")):
            p = os.path.join(base, rel)
            os.makedirs(os.path.dirname(p), exist_ok=True)
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = sheet
            ws.append(header)
            for r in rows:
                ws.append(r)
            wb.save(p)  # an .xlsm saved by openpyxl without VBA still loads read-only
            return p

        def folders(*rels):
            for r in rels:
                os.makedirs(os.path.join(tmp, r), exist_ok=True)

        book(root, "ita/italy_master.xlsm", [(1, 1, 216435, "Stanislav Lobotka", "real"), (2, 2, 216435, "Lobotka", "real"),
                                             (3, 3, 216435, "Lobotka", "real"), (4, 4, 920014, "Abbiati", "generic"),
                                             (5, 5, "n/a", "Broken", "real"), (6, 6, 244263.0, "Amir Rrahmani", "Real "),
                                             (None, None, None, None, None), (7, 7, 920014, "Abbiati", "generic"),
                                             (8, 8, 12, "Other", "unknown")])
        old = book(root, "ned/dutch_master.xlsx", [(1, 1, 1, "Old", "real")])
        new = book(root, "ned/dutch_master.xlsm", [(1, 1, 2, "New", "real")])
        os.utime(old, (2_000_000_000, 2_000_000_000))  # the .xlsx is newer, the .xlsm still wins
        book(root, "misc/spain_master.xlsm", [(1, 1, 3, "Tres", "real")])
        book(root, "eng/uk_master.xlsm", [(1, 1, 4, "x", "real")], sheet="other")
        book(root, "fra/france_master.xlsm", [(1, 1, 5, "x", "real")], header=("a", "b", "c"))
        # audio: ita has its own generic\ (FC 26), ned keeps its wavs in ned\audio\{generic,real} (as the user's ned)
        folders("My Mods/ita/generic", "My Mods/ned/audio/generic", "My Mods/ned/audio/Real")
        logs = []
        rc, res = run(root, out, log=logs.append)
        ok = True

        def check(cond, what):
            nonlocal ok
            print(("  ok    " if cond else "  FAIL  ") + what)
            ok = ok and cond

        def load(folder, lang):
            with open(os.path.join(folder, f"{lang}.json"), encoding="utf-8") as f:
                return json.load(f)

        check(rc == 0, "run succeeded")
        check(sorted(res) == ["dut_nl", "ita_it", "spa_es"], f"languages: {sorted(res)}")
        it = load(out, "ita_it")
        check(it["language"] == "ita_it" and it["real_players"] == [216435, 244263], f"real players: {it['real_players']}")
        check(it["game"] == "fc26" and it["note"].startswith("FC 26 data"), f"an FC 26 list says so: {it['game']}")
        check(it["generic_ids"] == [920014], f"generic ids: {it['generic_ids']}")
        check(it["names"] == {"216435": "Lobotka", "244263": "Amir Rrahmani"}, f"names (most frequent): {it['names']}")
        check(it["counts"]["non_numeric_id_rows"] == 1 and it["counts"]["other_type_rows"] == 1 and it["counts"]["blank_rows"] == 1,
              f"counts: {it['counts']}")
        check(list(it) == ["turbo_masters", "language", "game", "source", "built", "note", "counts", "real_players",
                           "generic_ids", "names", "generic_names", "wav_dir", "segments"]
              and list(it["segments"]) == ["generic", "real"] and "game_build" not in it,
              f"an FC 26 list keeps its shape (no FC 27 keys): {list(it)}")
        check(it["segments"] == {"generic": {"920014": [4, 7]}, "real": {"216435": [1, 2, 3], "244263": [6]}},
              f"segments: {it['segments']}")
        check(it["wav_dir"] == os.path.join(root, "ita"), f"FC 26 ita: the workbook's folder: {it.get('wav_dir')}")
        nl = load(out, "dut_nl")
        check(nl["source"] == os.path.abspath(new) and nl["real_players"] == [2], "dutch: the .xlsm wins over the newer .xlsx")
        check(nl.get("wav_dir") == os.path.join(root, "ned", "audio"), f"dutch: <workbook folder>\\audio: {nl.get('wav_dir')}")
        check("wav_dir" not in load(out, "spa_es"), "spa_es: no audio folder, no wav_dir")
        check(any("dutch_master.xlsx (FC 26 list): skipped" in s for s in logs), "the skipped duplicate is reported")
        check(any(s.startswith("eng_us:") and "NOT imported" in s and "no sheet" in s for s in logs), "missing sheet reported")
        check(any(s.startswith("fre_fr:") and "NOT imported" in s and "no column" in s for s in logs), "missing columns reported")
        check(not os.path.exists(os.path.join(out, "eng_us.json")), "nothing written for a failed workbook")
        rc2, _ = run(root, os.path.join(tmp, "dry"), dry_run=True, log=lambda s: None)
        check(rc2 == 0 and not os.path.exists(os.path.join(tmp, "dry")), "dry run writes nothing")
        # --wav-dir: one language's folder, '' = none
        outw = os.path.join(tmp, "outw")
        run(root, outw, langs={"dut_nl", "ita_it"}, log=lambda s: None,
            wav_dirs=parse_lang_map([f"dut_nl={os.path.join(tmp, 'elsewhere')}", "ita_it="], "--wav-dir"))
        check(load(outw, "dut_nl")["wav_dir"] == os.path.join(tmp, "elsewhere") and "wav_dir" not in load(outw, "ita_it"),
              "--wav-dir sets a folder (even one without audio, with a warning) and 'lang=' sets none")
        try:
            parse_lang_map(["nolang"], "--wav-dir")
            check(False, "a --wav-dir without '=' is refused")
        except ValueError:
            check(True, "a --wav-dir without '=' is refused")
        # an FC 27 master (same format, <stem>_master_fc27) in a second folder wins over the FC 26 list of its language,
        # whatever its extension or age; one in the same folder as the FC 26 list wins too
        fc27 = book(root27, "italy_master_fc27.xlsx", [(1, 1, 261865, "Miguel Gutierrez", "real"), (2, 2, 920014, "Abbiati", "generic")])
        os.utime(fc27, (1_000_000_000, 1_000_000_000))
        es27 = book(root, "misc/spain_master_fc27.xlsm", [(1, 1, 30, "Treinta", "real")])
        logs = []
        out27 = os.path.join(tmp, "out27")
        rc3, _ = run([root, root27, os.path.join(tmp, "missing")], out27, log=logs.append)
        it27 = load(out27, "ita_it")
        check(rc3 == 0 and it27["game"] == "fc27" and it27["source"] == os.path.abspath(fc27) and it27["real_players"] == [261865],
              f"ita_it from the FC 27 master: {it27['game']} {it27['real_players']}")
        check(it27["note"].startswith("FC 27 master"), "the note says FC 27")
        check("real_link_players" not in it27 and "game_build" not in it27,
              "an FC 27 master without FC 27 types or build: no FC 27 keys")
        check(any("italy_master.xlsm (FC 26 list): skipped" in s and "FC 27 master first" in s for s in logs),
              "the FC 26 list is reported as skipped")
        es = load(out27, "spa_es")
        check(es["game"] == "fc27" and es["source"] == os.path.abspath(es27), "spa_es: the FC 27 master next to the FC 26 list wins")
        check(load(out27, "dut_nl")["game"] == "fc26", "dut_nl: FC 26 list (no FC 27 master)")
        check(any("no folder" in s and "missing" in s for s in logs), "a missing folder is noted, not fatal")
        check(workbook_kind("italy_master_fc27.xlsm") == ("italy", "fc27") and workbook_kind("italy_master.xlsx") == ("italy", "fc26")
              and workbook_kind("~$italy_master.xlsm") is None and workbook_kind("italy_master.csv") is None
              and workbook_kind("italy_masters.xlsm") is None, "file name rules")

        # FC 27 types (a tiny .xlsx as build_callname_master.py writes it): 100 SIMPLE + LINK, 500 LINK only (its name from
        # the LINK row), 100 / 600 HIGH (not own recordings), 900001 generic + generic HIGH; the i27 audio folder; the
        # game build from the JSON build_callname_master wrote for this workbook, then from the bank
        root3 = os.path.join(tmp, "fc27")
        hdr = ("SegmentID", "VariationId", "playerid", "name", "type", "Play Audio", "Category", "nameid")
        x27 = book(root3, "ita/italy_master_fc27.xlsx", [
            (1, 11, 100, "Zola", "real", None, "Italian", None), (2, 12, 100, "Zola", "real", None, None, None),
            (5, 15, 200, "Rossi", "real", None, None, None), (0, 10, 200, "Rossi", "real", None, None, None),
            (9, 19, 100, "Gianfranco Zola", "real link", None, None, None), (10, 20, 500, "Linkonly", "Real Link", None, None, None),
            (10, 21, 500, "Linkonly", "real_link", None, None, None),
            (11, 31, 100, "Zola", "real high", None, None, None), (13, 33, 600, "Highonly", "REAL HIGH", None, None, None),
            (3, 13, 900001, "Rossi", "generic", None, None, "6"), (12, 32, 900001, "Rossi", "generic high", None, None, None),
            (14, 34, 970001, "Big", "generic", None, None, None), (15, 35, 7, "x", "real low", None, None, None)], header=hdr)
        folders("fc27/i27/generic", "fc27/i27/real", "fc27/i27/real_link", "fc27/ita/generic")
        with open(os.path.join(root3, "ita_it.json"), "w", encoding="utf-8") as f:   # as build_callname_master --out-json
            json.dump({"language": "ita_it", "game": "fc27", "game_build": "1.0.test",
                       "source": "C:\\elsewhere\\ita\\italy_master_fc27.xlsx"}, f)
        os.makedirs(os.path.join(root3, "raw"))
        with open(os.path.join(root3, "raw", "ita_it_bank.json"), "w", encoding="utf-8") as f:
            json.dump({"turbo_bank": 1, "lang": "ita_it", "game_build": "1.0.bank"}, f)
        logs = []
        out3 = os.path.join(tmp, "out3")
        rc4, _ = run(root3, out3, log=logs.append)
        j = load(out3, "ita_it")
        check(rc4 == 0 and j["game"] == "fc27" and j["source"] == os.path.abspath(x27), "FC 27 types: imported")
        check(j["real_players"] == [100, 200, 500] and j["real_simple_players"] == [100, 200] and j["real_link_players"] == [100, 500],
              f"real_players = real | real link: {j['real_players']} {j.get('real_simple_players')} {j.get('real_link_players')}")
        check(j["real_high_players"] == [100, 600] and j["generic_high_ids"] == [900001] and j["generic_ids"] == [900001, 970001],
              f"HIGH lists (not own recordings): {j.get('real_high_players')} {j.get('generic_high_ids')}")
        check(j["names"] == {"100": "Zola", "200": "Rossi", "500": "Linkonly"}, f"names: 'real' first, else 'real link': {j['names']}")
        check(j["segments"] == {"generic": {"900001": [3], "970001": [14]}, "real": {"100": [1, 2], "200": [0, 5]},
                                "real_link": {"100": [9], "500": [10]}, "real_high": {"100": [11], "600": [13]},
                                "generic_high": {"900001": [12]}}, f"segments per family (SegmentID 0 kept): {j['segments']}")
        k = j["counts"]
        check(k["real_rows"] == 4 and k["real_link_rows"] == 3 and k["real_high_rows"] == 2 and k["generic_high_rows"] == 1
              and k["generic_rows"] == 2 and k["other_type_rows"] == 1 and k["real_players"] == 2 and k["real_link_players"] == 2
              and k["own_recording_players"] == 3 and k["real_high_players"] == 2 and k["generic_high_ids"] == 1
              and k["any_recording_players"] == 4 and k["generic_ids_outside_900001_965000"] == 1, f"counts: {k}")
        check(j.get("game_build") == "1.0.test" and list(j)[:5] == ["turbo_masters", "language", "game", "game_build", "source"],
              f"game_build from build_callname_master's JSON for the same workbook: {j.get('game_build')}")
        check(j.get("wav_dir") == os.path.join(root3, "i27"), f"FC 27 Italian: the i27 audio folder: {j.get('wav_dir')}")
        check("'real link'" in j["note"] and any("3 players with their own recording (2 'real', 4 rows | 2 'real link', 3 rows)"
                                                 in s for s in logs), "note and log name the FC 27 types")
        check(any(s.startswith("ita_it: wav_dir ") and "game_build 1.0.test" in s for s in logs), "audio folder and build logged")
        with open(os.path.join(root3, "ita_it.json"), "w", encoding="utf-8") as f:   # a JSON of another workbook: not used
            json.dump({"language": "ita_it", "game": "fc27", "game_build": "1.0.test", "source": "x\\other_master_fc27.xlsm"}, f)
        _, r5 = run(root3, None, dry_run=True, log=lambda s: None)
        check(r5["ita_it"].get("game_build") == "1.0.bank", f"game_build from the bank: {r5['ita_it'].get('game_build')}")
        _, r6 = run(root3, None, dry_run=True, log=lambda s: None, game_builds=parse_lang_map(["ita_it=9.9"], "--game-build"))
        check(r6["ita_it"].get("game_build") == "9.9", "--game-build wins")
        print("SELF-TEST " + ("PASSED" if ok else "FAILED"))
        return 0 if ok else 1


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--root", action="append", help="folder with the language folders or the workbooks; repeatable "
                    "(default: C:\\FC_Tools\\My Mods and <le>\\turbo_dev\\masters, each when present)")
    ap.add_argument("--le", default=default_le_root(), help="Live Editor folder (default: %(default)s, or TURBO_LE_ROOT)")
    ap.add_argument("--out", help="output folder (default: <le>\\turbo\\callnames\\masters)")
    ap.add_argument("--lang", nargs="*", help="only these language codes (e.g. ita_it)")
    ap.add_argument("--dry-run", action="store_true", help="read and report, write nothing")
    ap.add_argument("--wav-dir", action="append", metavar="LANG=DIR",
                    help="audio folder of one language (holding generic\\ and real\\); LANG= writes none; repeatable")
    ap.add_argument("--game-build", action="append", metavar="LANG=BUILD",
                    help="game build of an FC 27 master (default: from build_callname_master's JSON or the bank)")
    ap.add_argument("--self-test", action="store_true", help="run the built-in test on synthetic workbooks and exit")
    a = ap.parse_args()
    if a.self_test:
        return self_test()
    try:
        wav_dirs = parse_lang_map(a.wav_dir, "--wav-dir")
        game_builds = parse_lang_map(a.game_build, "--game-build")
    except ValueError as e:
        ap.error(str(e))
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        print("error: openpyxl is not installed (python -m pip install openpyxl)", file=sys.stderr)
        return 2
    out = a.out or os.path.join(a.le, "turbo", "callnames", "masters")
    rc, res = run(a.root or default_roots(a.le), out, set(x.lower() for x in a.lang) if a.lang else None, a.dry_run,
                  wav_dirs=wav_dirs, game_builds=game_builds)
    if res:
        print(f"{len(res)} language list(s) {'read' if a.dry_run else 'written to ' + out}; restart the Turbo window's "
              f"Callname tab (Refresh) to load them")
    return rc


if __name__ == "__main__":
    sys.exit(main())
